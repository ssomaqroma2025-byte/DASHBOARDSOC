# -*- coding: utf-8 -*-
# Convierte los Excel de PERSONAL y Justificaciones a los CSV que usa el robot.
# Uso: python app/excel2csv.py <carpeta_con_excel>   (por defecto data/excel)
# Reconoce los archivos por su nombre: "PERSONAL*.xlsx" y "Justificaciones*.xlsx".
import os, sys, csv, glob, io
import openpyxl
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from respuesta import motivo_publico   # el motivo real (DM, etc.) no se publica

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(ROOT, "data")
SRC  = sys.argv[1] if len(sys.argv) > 1 else os.path.join(DATA, "excel")

def norm(s): return "" if s is None else str(s).strip()

def header_rows(ws, must, maxscan=5):
    """Devuelve (indices, iterador) buscando la fila de encabezado que contiene 'must'."""
    it = ws.iter_rows(values_only=True)
    for _ in range(maxscan):
        h = [norm(x) for x in next(it)]
        if must in h:
            return {v: i for i, v in enumerate(h) if v}, it
    raise ValueError("No encontré el encabezado %r en la hoja %s" % (must, ws.title))

def write_csv(name, header, rows):
    p = os.path.join(DATA, name)
    with io.open(p, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f); w.writerow(header); w.writerows(rows)
    print("  %s: %d filas" % (name, len(rows)))

def personal(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ix, it = header_rows(wb["Usuarios_operarios"], "Codigo")
    get = lambda r, k: norm(r[ix[k]]) if k in ix and ix[k] < len(r) else ""
    plk = "Planta" if "Planta" in ix else "Planta "
    rows, seen = [], set()
    for r in it:
        cod = get(r, "Codigo")
        if not cod or cod in seen: continue
        seen.add(cod)
        rows.append([cod, get(r, plk), get(r, "Gerencia"), get(r, "Area"),
                     get(r, "Apellidos y Nombres"), get(r, "Empleado")])
    write_csv("personal.csv", ["cod","planta","gerencia","area","nombre","tipo"], rows)
    # hoja de bajas/altas/DM dentro de PERSONAL
    ix, it = header_rows(wb["DM VACACIONES BAJAS Y ALTAS"], "CÓDIGO")
    rows = [[norm(r[ix["CÓDIGO"]]), norm(r[ix["AÑO"]]), norm(r[ix["SEMANA"]]), motivo_publico(norm(r[ix["TIPO"]]), "b")]
            for r in it if norm(r[ix["CÓDIGO"]])]
    write_csv("justif_b.csv", ["cod","anio","semana","motivo"], rows)

def justificaciones(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ix, it = header_rows(wb[wb.sheetnames[0]], "CÓDIGO (NECESARIO)")   # solo la primera hoja
    rows = [[norm(r[ix["CÓDIGO (NECESARIO)"]]), norm(r[ix["AÑO"]]), norm(r[ix["SEMANA"]]), "J"]
            for r in it if norm(r[ix["CÓDIGO (NECESARIO)"]])]
    write_csv("justif_a.csv", ["cod","anio","semana","motivo"], rows)

if __name__ == "__main__":
    hechos = 0
    for f in glob.glob(os.path.join(SRC, "*.xlsx")):
        n = os.path.basename(f).upper()
        if n.startswith("~$"): continue
        if n.startswith("PERSONAL"):          print("PERSONAL:", f); personal(f); hechos += 1
        elif n.startswith("JUSTIFICACIONES"): print("Justificaciones:", f); justificaciones(f); hechos += 1
    print("Archivos convertidos:", hechos)
