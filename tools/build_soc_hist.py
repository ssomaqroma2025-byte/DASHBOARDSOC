# -*- coding: utf-8 -*-
# Lee los Excel de respuestas de los formularios (Agustino y Ñaña) y genera data/soc_hist.json:
#   soc[codigo][semanaISO] = cantidad de SOC, y "corte" = fecha-hora de la última respuesta leída.
# Uso: python tools/build_soc_hist.py "<Excel Agustino>" "<Excel Ñaña>"
import os, sys, io, json, re, datetime
from collections import defaultdict
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT  = os.path.join(ROOT, "data", "soc_hist.json")
MINYEAR = 2024

def norm(s): return "" if s is None else str(s).strip()

def to_dt(v):
    if isinstance(v, datetime.datetime): return v
    if isinstance(v, datetime.date): return datetime.datetime(v.year, v.month, v.day)
    m = re.match(r"(\d{4})-(\d{2})-(\d{2})(?:[ T](\d{2}):(\d{2})(?::(\d{2}))?)?", norm(v))
    if not m: return None
    y, mo, d, h, mi, s = (int(x) if x else 0 for x in m.groups())
    return datetime.datetime(y, mo, d, h, mi, s)

soc = defaultdict(lambda: defaultdict(int))
corte = None
for path in sys.argv[1:]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["Sheet1"] if "Sheet1" in wb.sheetnames else wb.worksheets[0]
    it = ws.iter_rows(min_col=1, max_col=12, values_only=True)
    h = [norm(x) for x in next(it)]
    iF, iC = h.index("Hora de inicio"), h.index("CÓDIGO DEL OBSERVADOR")
    iE = h.index("Hora de finalización") if "Hora de finalización" in h else iF
    n = 0
    for r in it:
        d = to_dt(r[iF]); cod = norm(r[iC])
        if not d or not cod or d.year < MINYEAR: continue
        iy, iw, _ = d.date().isocalendar()
        soc[cod][str(iy*100+iw)] += 1; n += 1
        e = to_dt(r[iE]) or d
        if corte is None or e > corte: corte = e
    print("%s: %d respuestas" % (os.path.basename(path), n))

data = {"corte": corte.strftime("%Y-%m-%dT%H:%M:%S") if corte else None,
        "soc": {c: dict(v) for c, v in soc.items()}}
io.open(OUT, "w", encoding="utf-8").write(json.dumps(data, ensure_ascii=False))
print("soc_hist.json: %d códigos, corte %s, %d KB" % (len(data["soc"]), data["corte"], os.path.getsize(OUT)//1024))
