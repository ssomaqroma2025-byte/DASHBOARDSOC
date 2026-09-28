# -*- coding: utf-8 -*-
# Arma data/det_hist.json (detalle de cada SOC para Análisis / Comentarios / Detalles)
# y data/areas.json (catálogo AREAS Y SUBAREAS: gerencia + títulos de C1..C4)
# a partir de los Excel descargados de los formularios.
#   python tools/build_det.py "<SOC Agustino.xlsx>" "<SOC Ñaña y Provincia.xlsx>" "<AREAS Y SUBAREAS SOC.xlsx>"
# Misma lógica que la consulta "SOC AGUS" del Power BI:
#   ÁREA = área de la planta (o provincia), SUBÁREA = la subárea marcada (si no hay, el área),
#   C1..C4 = la respuesta de la columna "comportamiento 0k" que tenga dato; vacío = NO SE OBSERVÓ.
import sys, os, io, re, json, datetime, unicodedata
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(ROOT, "data")

def norm(s):
    if s is None: return ""
    return re.sub(r"\s+", " ", str(s).replace("\xa0", " ")).strip()
def key(s):
    s = unicodedata.normalize("NFD", norm(s).upper())
    return "".join(c for c in s if unicodedata.category(c) != "Mn")

VAL = {"SEGURO": 0, "INSEGURO": 1}          # otro / vacío = 2 (NO SE OBSERVÓ)
BASURA = {"", ".", "..", "...", "0", ",", "-"}

def planta_de(v):
    k = key(v)
    if k.startswith("AGUSTINO"): return "AGUSTINO"
    if k.startswith("NANA"): return "ÑAÑA"
    if k: return "PROVINCIAS"
    return ""

out = dict(ub=[], w=[], u=[], t=[], c=[], k=[], m=[], com=[])
ubi = {}
vocab_areas, vocab_subs, vocab_clasif = {}, set(), set()
corte = ""

def leer(path):
    global corte
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["Sheet1"] if "Sheet1" in wb.sheetnames else wb[wb.sheetnames[0]]
    it = ws.iter_rows(min_col=1, max_col=260, values_only=True)
    hdr = [norm(h) for h in next(it)]
    ix = {h: i for i, h in enumerate(hdr) if h}
    def col(pref):
        for h, i in ix.items():
            if key(h).startswith(key(pref)): return i
        return None
    iF, iFin = ix["Hora de inicio"], ix.get("Hora de finalización")
    iC = ix["CÓDIGO DEL OBSERVADOR"]
    iT = col("SELECCIONA EL TURNO"); iP = col("SELECCIONA LA PLANTA")
    iAg = col("SELECCIONA EL ÁREA OBSERVADA - AGUSTINO"); iNa = col("SELECCIONA EL ÁREA OBSERVADA - ÑAÑA")
    iPr = col("SELECCIONE LA PROVINCIA")
    iCl = col("Clasifica la observación"); iDe = col("Detalla lo observado")
    subs = [i for h, i in ix.items() if key(h).startswith("SUBAREA")]
    comp = {1: [], 2: [], 3: [], 4: []}
    for h, i in ix.items():
        mm = re.search(r"comportamiento\s*0?([1-4])", h, re.I)
        if mm: comp[int(mm.group(1))].append(i)
    n = 0
    for r in it:
        f = r[iF]
        if not isinstance(f, datetime.datetime): continue
        cod = norm(r[iC])
        if cod.endswith(".0"): cod = cod[:-2]
        if not cod: continue
        pl = planta_de(r[iP] if iP is not None else "")
        if not pl: continue
        if pl == "AGUSTINO": ar = norm(r[iAg]) if iAg is not None else ""
        elif pl == "ÑAÑA":   ar = norm(r[iNa]) if iNa is not None else ""
        else:                ar = norm(r[iPr]) if iPr is not None else ""
        ar = ar or "(sin área)"
        sub = next((norm(r[i]) for i in subs if norm(r[i])), "") or ar
        tt = re.search(r"(\d)", norm(r[iT]) if iT is not None else "")
        kk = 0
        for p in (1, 2, 3, 4):
            v = next((norm(r[i]) for i in comp[p] if norm(r[i])), "")
            kk += VAL.get(key(v), 2) * 3 ** (p - 1)
        de = norm(r[iDe]) if iDe is not None else ""
        cl = norm(r[iCl]) if iCl is not None else ""
        if cl: vocab_clasif.add(cl)
        vocab_areas.setdefault(pl, set()).add(ar); vocab_subs.add(sub)
        uk = (pl, ar, sub)
        if uk not in ubi: ubi[uk] = len(out["ub"]); out["ub"].append(list(uk))
        iy, iw, _ = f.date().isocalendar()
        out["w"].append(iy * 100 + iw); out["u"].append(ubi[uk]); out["t"].append(int(tt.group(1)) if tt else 0)
        out["c"].append(cod); out["k"].append(kk)
        if de not in BASURA:
            out["m"].append(len(out["com"])); out["com"].append(de)
        else:
            out["m"].append(-1)
        fin = r[iFin] if iFin is not None else None
        if isinstance(fin, datetime.datetime): corte = max(corte, fin.strftime("%Y-%m-%dT%H:%M:%S"))
        n += 1
    wb.close()
    print("  %s: %d respuestas" % (os.path.basename(path), n))

def catalogo(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True); hdr = [norm(h) for h in next(it)]
    ix = {h: i for i, h in enumerate(hdr)}
    def tcols(suf): return [ix.get("TITULO COMPORTAMIENTO %d%s" % (p, suf)) for p in (1, 2, 3, 4)]
    vers = [tcols(""), tcols("_ANTIGUO_1"), tcols("_ANTIGUO_2"), tcols("_ANTIGUO_3"), tcols("_ANTIGUO_4")]
    rows = []
    for r in it:
        if not r or not norm(r[ix["PLANTA"]]): continue
        rows.append(dict(pl=norm(r[ix["PLANTA"]]), ar=norm(r[ix["AREA"]]), sub=norm(r[ix["SUBAREA"]]),
                         ger=norm(r[ix["GERENCIA"]]),
                         t=[[norm(r[i]) if i is not None else "" for i in v] for v in vers]))
    wb.close()
    return rows

if __name__ == "__main__":
    agus, nana, areas = sys.argv[1:4]
    for p in (agus, nana): leer(p)
    out["corte"] = corte
    out["vocab"] = dict(areas={k: sorted(v) for k, v in vocab_areas.items()},
                        subs=sorted(vocab_subs), clasif=sorted(vocab_clasif))
    json.dump(out, io.open(os.path.join(DATA, "det_hist.json"), "w", encoding="utf-8"),
              ensure_ascii=False, separators=(",", ":"))
    json.dump(catalogo(areas), io.open(os.path.join(DATA, "areas.json"), "w", encoding="utf-8"),
              ensure_ascii=False, separators=(",", ":"))
    print("det_hist.json: %d SOC, %d comentarios, %d ubicaciones, corte %s" %
          (len(out["w"]), len(out["com"]), len(out["ub"]), corte))
